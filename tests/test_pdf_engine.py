import os
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pypdf import PdfReader
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from pdf_engine import (
    extract_pdf_pages,
    extract_pdf_text,
    images_to_pdf,
    merge_pdfs,
    pdf_to_docx,
    pdf_to_pptx,
    pdf_to_xlsx,
    render_pdf_pages,
    rotate_pdf,
    split_pdf,
)


class PdfEngineTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.first = self.root / "first.pdf"
        self.second = self.root / "second.pdf"
        Image.new("RGB", (120, 80), "red").save(self.first, "PDF")
        Image.new("RGB", (160, 100), "blue").save(self.second, "PDF")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_merge_split_rotate_and_extract(self):
        merged = self.root / "merged.pdf"
        merge_pdfs([str(self.first), str(self.second)], str(merged))
        self.assertEqual(len(PdfReader(str(merged)).pages), 2)

        pages = split_pdf(str(merged), str(self.root))
        self.assertEqual(len(pages), 2)
        self.assertTrue(all(Path(path).is_file() for path in pages))

        rotated = self.root / "rotated.pdf"
        rotate_pdf(str(merged), str(rotated), 90)
        self.assertEqual(len(PdfReader(str(rotated)).pages), 2)

        extracted = self.root / "extracted.pdf"
        extract_pdf_pages(str(merged), str(extracted), [2])
        self.assertEqual(len(PdfReader(str(extracted)).pages), 1)

    def test_render_and_images_to_pdf(self):
        rendered = render_pdf_pages(str(self.first), str(self.root), "PNG")
        self.assertEqual(len(rendered), 1)
        self.assertTrue(Path(rendered[0]).is_file())

        jpeg = render_pdf_pages(str(self.first), str(self.root), "JPEG")
        self.assertEqual(len(jpeg), 1)
        self.assertEqual(Path(jpeg[0]).suffix, ".jpg")

        output = self.root / "images.pdf"
        images_to_pdf(rendered, str(output))
        self.assertEqual(len(PdfReader(str(output)).pages), 1)

    def test_text_and_best_effort_office_exports(self):
        source = self.root / "content.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(612, 792)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref}),
        })
        content = (
            b"0.5 w 72 600 m 272 600 l S 72 650 m 272 650 l S "
            b"72 700 m 272 700 l S 72 600 m 72 700 l S "
            b"172 600 m 172 700 l S 272 600 m 272 700 l S "
            b"BT /F1 12 Tf 90 670 Td (Name) Tj 100 0 Td (Value) Tj "
            b"-100 -50 Td (One) Tj 100 0 Td (1) Tj ET"
        )
        stream = DecodedStreamObject()
        stream.set_data(content)
        page[NameObject("/Contents")] = writer._add_object(stream)
        with open(source, "wb") as output:
            writer.write(output)

        text_path = self.root / "content.txt"
        self.assertIn("Name", extract_pdf_text(str(source), str(text_path)))

        docx_path = self.root / "content.docx"
        pdf_to_docx(str(source), str(docx_path))
        self.assertIn("Name", "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs))

        xlsx_path = self.root / "content.xlsx"
        pdf_to_xlsx(str(source), str(xlsx_path))
        self.assertEqual(load_workbook(xlsx_path).active["A1"].value, "Name")

        pptx_path = self.root / "content.pptx"
        pdf_to_pptx(str(source), str(pptx_path))
        self.assertEqual(len(Presentation(pptx_path).slides), 1)


if __name__ == "__main__":
    unittest.main()
